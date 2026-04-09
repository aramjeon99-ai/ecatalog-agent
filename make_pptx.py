from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

bg = slide.background; bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(0xF0, 0xF2, 0xF5)

I = Inches
RNDRT = 5   # rounded rectangle

# ── helpers ──────────────────────────────────────────────────────────
def box(l, t, w, h, fill, brdr, lines, bw=Pt(1.5), stype=RNDRT):
    s = slide.shapes.add_shape(stype, I(l), I(t), I(w), I(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(*fill)
    s.line.color.rgb = RGBColor(*brdr)
    s.line.width = bw
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (txt_str, sz, bold, clr) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(1)
        p.space_after  = Pt(1)
        run = p.add_run()
        run.text = txt_str
        run.font.size = Pt(sz)
        run.font.bold = bold
        if clr:
            run.font.color.rgb = RGBColor(*clr)
    return s

def txt(l, t, w, h, text, sz=10, bold=False, clr=(0x44,0x44,0x44)):
    tb = slide.shapes.add_textbox(I(l), I(t), I(w), I(h))
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(sz)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*clr)
    return tb

def add_arrow_v(cx, y1, y2):
    conn = slide.shapes.add_connector(1, I(cx), I(y1), I(cx), I(y2))
    conn.line.color.rgb = RGBColor(0x88, 0x88, 0x88)
    conn.line.width = Pt(1.5)
    sp = conn._element.spPr
    ln = sp.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(sp, qn('a:ln'))
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'arrow'); tail.set('w', 'med'); tail.set('len', 'med')

def add_arrow_h(y, x1, x2):
    conn = slide.shapes.add_connector(1, I(x1), I(y), I(x2), I(y))
    conn.line.color.rgb = RGBColor(0x88, 0x88, 0x88)
    conn.line.width = Pt(1.5)
    sp = conn._element.spPr
    ln = sp.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(sp, qn('a:ln'))
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'arrow'); tail.set('w', 'med'); tail.set('len', 'med')

def group_rect(l, t, w, h, brdr_clr, fill_clr=(0xFF,0xFF,0xFF), alpha=180):
    s = slide.shapes.add_shape(1, I(l), I(t), I(w), I(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(*fill_clr)
    s.line.color.rgb = RGBColor(*brdr_clr)
    s.line.width = Pt(1.2)
    sp = s._element.spPr
    ln = sp.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(sp, qn('a:ln'))
    pd = etree.SubElement(ln, qn('a:prstDash'))
    pd.set('val', 'dash')
    return s

# ══════════════════════════════════════════════════════════════════════
# 제목
txt(0.3, 0.05, 12.73, 0.42,
    'eCatalog Agent  —  시스템 아키텍처',
    sz=18, bold=True, clr=(0x1A,0x1A,0x2E))

# ── ROW 1 : Streamlit UI ─────────────────────────────────────────────
group_rect(0.3, 0.55, 6.3, 1.0,
           (0x34,0x98,0xDB), (0xEB,0xF5,0xFB))
txt(0.48, 0.57, 3.0, 0.22,
    'Streamlit UI', sz=8, bold=True, clr=(0x1A,0x6F,0xA8))

box(0.5, 0.68, 2.8, 0.78,
    (0xD6,0xEA,0xF8), (0x34,0x98,0xDB),
    [('INPUT', 8, True, (0x1A,0x5E,0x7A)),
     ('시스템 데이터 업로드', 9, True, (0x1A,0x1A,0x2E)),
     ('system_data / pdf_mapping', 7, False, (0x55,0x55,0x55))])

add_arrow_h(1.07, 3.35, 3.85)

box(3.9, 0.68, 2.5, 0.78,
    (0xD6,0xEA,0xF8), (0x34,0x98,0xDB),
    [('LIST', 8, True, (0x1A,0x5E,0x7A)),
     ('Q코드 목록', 9, True, (0x1A,0x1A,0x2E)),
     ('PDF·URL 있는 항목  |  앞7+뒤7', 7, False, (0x55,0x55,0x55))])

# ── Arrow ────────────────────────────────────────────────────────────
add_arrow_v(6.665, 1.58, 1.78)

# ── ROW 2 : QcodeContext ─────────────────────────────────────────────
box(4.33, 1.82, 4.67, 0.78,
    (0xD1,0xF2,0xEB), (0x1A,0xBC,0x9C),
    [('CONTEXT', 8, True, (0x0E,0x6B,0x5B)),
     ('QcodeContext 수집', 9, True, (0x1A,0x1A,0x2E)),
     ('PDF 텍스트(PyMuPDF)  ·  URL  ·  사양값  ·  메이커명', 7, False, (0x44,0x44,0x44))])

# ── Arrow ────────────────────────────────────────────────────────────
add_arrow_v(6.665, 2.65, 2.82)

# ── ROW 3 : 병렬 실행 ────────────────────────────────────────────────
group_rect(0.3, 2.85, 12.73, 0.92,
           (0xF3,0x9C,0x12), (0xFF,0xFD,0xF0))
txt(0.48, 2.87, 5.0, 0.22,
    'ThreadPoolExecutor  (병렬 실행)',
    sz=8, bold=True, clr=(0xB7,0x77,0x0E))

box(0.5, 3.03, 3.9, 0.65,
    (0xFE,0xF9,0xE7), (0xF3,0x9C,0x12),
    [('메이커 검증', 9, True, (0x1A,0x1A,0x2E)),
     ('verify_manufacturer()  |  maker_catalog_hints', 7, False, (0x66,0x55,0x44))])

box(4.72, 3.03, 3.9, 0.65,
    (0xFE,0xF9,0xE7), (0xF3,0x9C,0x12),
    [('웹 검색', 9, True, (0x1A,0x1A,0x2E)),
     ('web_searcher  (timeout 6s)', 7, False, (0x66,0x55,0x44))])

box(8.93, 3.03, 3.9, 0.65,
    (0xFE,0xF9,0xE7), (0xF3,0x9C,0x12),
    [('URL 스펙 추출', 9, True, (0x1A,0x1A,0x2E)),
     ('url_spec_fetcher  (timeout 8s)', 7, False, (0x66,0x55,0x44))])

# ── Arrow ────────────────────────────────────────────────────────────
add_arrow_v(6.665, 3.82, 3.99)

# ── ROW 4 : 조기 승인 로직 ──────────────────────────────────────────
group_rect(0.3, 4.02, 12.73, 1.55,
           (0x27,0xAE,0x60), (0xEA,0xF9,0xF1))
txt(0.48, 4.04, 4.0, 0.22,
    'Early-Exit  (조기 승인 로직)',
    sz=8, bold=True, clr=(0x1A,0x6B,0x40))

# 왼쪽 — _pdf_fully_confirmed
box(0.5, 4.22, 5.8, 0.62,
    (0xD5,0xF5,0xE3), (0x27,0xAE,0x60),
    [('_pdf_fully_confirmed', 9, True, (0x1A,0x4A,0x30)),
     ('pdf존재 + 모델일치 + 메이커일치 + 사양확인', 7, False, (0x33,0x55,0x44))])

txt(0.85, 4.90, 1.1, 0.22, 'YES', sz=8, bold=True, clr=(0x27,0xAE,0x60))
add_arrow_h(5.01, 1.55, 2.3)
box(2.35, 4.89, 2.5, 0.46,
    (0x27,0xAE,0x60), (0x1E,0x8A,0x49),
    [('즉시 승인  ✅', 9, True, (0xFF,0xFF,0xFF))])

txt(4.95, 4.90, 0.9, 0.22, 'NO', sz=8, bold=True, clr=(0xE7,0x4C,0x3C))
add_arrow_h(5.01, 5.55, 6.0)
box(6.03, 4.89, 0.85, 0.46,
    (0xD5,0xD8,0xDC), (0x95,0xA5,0xA6),
    [('다음', 8, False, (0x44,0x44,0x44))])

# 오른쪽 — _no_meaningful_specs
box(6.95, 4.22, 5.85, 0.62,
    (0xD5,0xF5,0xE3), (0x27,0xAE,0x60),
    [('_no_meaningful_specs', 9, True, (0x1A,0x4A,0x30)),
     ('사양값 전부 9999 (미입력 상태)', 7, False, (0x33,0x55,0x44))])

txt(7.3, 4.90, 1.1, 0.22, 'YES', sz=8, bold=True, clr=(0x27,0xAE,0x60))
add_arrow_h(5.01, 7.95, 8.7)
box(8.73, 4.89, 2.7, 0.46,
    (0x27,0xAE,0x60), (0x1E,0x8A,0x49),
    [('모델·메이커 일치 시 승인  ✅', 8, True, (0xFF,0xFF,0xFF))])

txt(11.5, 4.90, 0.9, 0.22, 'NO', sz=8, bold=True, clr=(0xE7,0x4C,0x3C))
add_arrow_h(5.01, 12.1, 12.55)
box(12.58, 4.89, 0.85, 0.46,
    (0xD5,0xD8,0xDC), (0x95,0xA5,0xA6),
    [('Step', 8, False, (0x44,0x44,0x44))])

# ── Arrow ────────────────────────────────────────────────────────────
add_arrow_v(6.665, 5.62, 5.77)

# ── ROW 5 : LangGraph ────────────────────────────────────────────────
group_rect(0.3, 5.8, 12.73, 1.08,
           (0x2C,0x3E,0x50), (0xEA,0xF0,0xFB))
txt(0.48, 5.82, 6.0, 0.22,
    'LangGraph 워크플로우  (graph.py)',
    sz=8, bold=True, clr=(0x2C,0x3E,0x50))

steps = [
    ('STEP 0', '입력 검증',   '필드·PDF 존재',     (0xEA,0xF0,0xFB), (0x2C,0x3E,0x50)),
    ('STEP 1', 'PDF 파싱',    '모델·메이커 매칭',   (0xEA,0xF0,0xFB), (0x2C,0x3E,0x50)),
    ('STEP 2', '신뢰도',      'PDF 출처 품질',      (0xEA,0xF0,0xFB), (0x2C,0x3E,0x50)),
    ('STEP 3', '사양 비교',   'Claude API',         (0xEA,0xF0,0xFB), (0x2C,0x3E,0x50)),
    ('STEP 4', '메이커 검증', 'maker_list + hints', (0xEA,0xF0,0xFB), (0x2C,0x3E,0x50)),
    ('STEP 5', '중복 검사',   'SQLite',             (0xEA,0xF0,0xFB), (0x2C,0x3E,0x50)),
    ('STEP 6', '최종 판정',   'APPROVED / REJECTED',(0xF4,0xEC,0xF7), (0x8E,0x44,0xAD)),
]
bw = 1.6; gap = 0.22; sx = 0.45
for i, (badge, title, sub, fill, brdr) in enumerate(steps):
    bx = sx + i * (bw + gap)
    tc = (0x2C,0x3E,0x50) if brdr != (0x8E,0x44,0xAD) else (0x8E,0x44,0xAD)
    box(bx, 5.97, bw, 0.82, fill, brdr,
        [(badge, 7, True, tc),
         (title, 9, True, (0x1A,0x1A,0x2E)),
         (sub,   7, False,(0x55,0x55,0x55))])
    if i < 6:
        add_arrow_h(6.38, bx + bw, bx + bw + gap)

# ── Arrow ────────────────────────────────────────────────────────────
add_arrow_v(6.665, 6.93, 7.07)

# ── ROW 6 : 출력 ─────────────────────────────────────────────────────
out_items = [
    ('검토 보고서',  '사양확장표  (.xlsx)',    (0xF2,0xF3,0xF4), (0x95,0xA5,0xA6)),
    ('SQLite 로그',  '단계별 판정 이력',       (0xF2,0xF3,0xF4), (0x95,0xA5,0xA6)),
    ('UI 결과 표시', '판정 · 이유 · 사양 비교표',(0xD6,0xEA,0xF8),(0x34,0x98,0xDB)),
]
ow = 3.6; ogap = 0.5
ox = (13.33 - 3*ow - 2*ogap) / 2
for i, (title, sub, fill, brdr) in enumerate(out_items):
    bx = ox + i * (ow + ogap)
    box(bx, 7.1, ow, 0.32, fill, brdr,
        [(title, 9, True, (0x1A,0x1A,0x2E)),
         (sub,   7, False,(0x55,0x55,0x55))])
    if i < 2:
        add_arrow_h(7.26, bx + ow, bx + ow + ogap)

prs.save('eCatalog_Agent_Architecture.pptx')
print('완료')
